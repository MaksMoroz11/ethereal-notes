"""Shared, restricted document tree for PDF and DOCX export."""
from dataclasses import dataclass, field
from html import escape
from html.parser import HTMLParser
from io import BytesIO
from urllib.parse import urlsplit

ALLOWED = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li", "blockquote",
           "strong", "b", "em", "i", "s", "strike", "u", "a", "code", "pre", "br", "hr"}
BLOCKED = {"script", "style", "iframe", "object", "svg", "math"}


@dataclass
class Node:
    tag: str
    attrs: dict = field(default_factory=dict)
    children: list = field(default_factory=list)


class DocumentParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.stack = [self.root]
        self.blocked = []

    def handle_starttag(self, tag, attrs):
        if self.blocked:
            if tag in BLOCKED:
                self.blocked.append(tag)
            return
        if tag in BLOCKED:
            self.blocked.append(tag)
            return
        if tag not in ALLOWED:
            return
        clean = {}
        for key, value in attrs:
            if key == "href" and tag == "a" and value and urlsplit(value).scheme in {"http", "https", "mailto"}:
                clean[key] = value
            if key == "start" and tag == "ol" and value and value.isdigit():
                clean[key] = str(min(int(value), 100000))
        node = Node(tag, clean)
        self.stack[-1].children.append(node)
        if tag not in {"br", "hr"}:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in {"br", "hr"}:
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        if self.blocked:
            if tag == self.blocked[-1]:
                self.blocked.pop()
            return
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index].tag == tag:
                del self.stack[index:]
                break

    def handle_data(self, data):
        if not self.blocked:
            self.stack[-1].children.append(data)


def parse_content(content):
    parser = DocumentParser()
    if "<" not in content:
        content = "".join(f"<p>{escape(line)}</p>" for line in content.splitlines())
    parser.feed(content)
    return parser.root


def html_of(node):
    if isinstance(node, str):
        return escape(node)
    children = "".join(html_of(child) for child in node.children)
    if node.tag == "root":
        return children
    attrs = "".join(f' {key}="{escape(value, quote=True)}"' for key, value in node.attrs.items())
    if node.tag == "ol" and "start" in node.attrs:
        attrs += f' style="counter-reset:list-item {int(node.attrs["start"]) - 1}"'
    return f"<{node.tag}{attrs}>{children}</{node.tag}>" if node.tag not in {"br", "hr"} else f"<{node.tag}>"


def generate_pdf(title, root):
    from weasyprint import HTML

    def deny_resource(url, **kwargs):
        raise ValueError("External resources are not supported")

    css = """@page {size:A4;margin:20mm;@bottom-center{content:counter(page);font-size:9pt}}
    body{font-family:'Noto Sans',sans-serif;font-size:11pt;line-height:1.5;overflow-wrap:anywhere}
    h1,h2,h3,h4,h5,h6{break-after:avoid} pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:10pt}
    blockquote{border-left:3px solid #aaa;padding-left:12px;color:#444} a{color:#395b88}
    """
    source = f"<!doctype html><html lang='ru'><meta charset='utf-8'><style>{css}</style><body><h1>{escape(title)}</h1>{html_of(root)}</body></html>"
    return HTML(string=source, url_fetcher=deny_resource).write_pdf()


def generate_docx(title, root):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(20)
    for name in ["Normal", "Title", "Quote", "List Number", "List Bullet", *[f"Heading {level}" for level in range(1, 7)]]:
        style = document.styles[name]
        style.font.name = "Noto Sans"
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        for key in list(fonts.attrib):
            if key.endswith("Theme"):
                del fonts.attrib[key]
        props = style.element.find(qn("w:pPr"))
        if props is not None:
            for border in props.findall(qn("w:pBdr")):
                props.remove(border)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.line_spacing = 1.3
    document.styles["Title"].font.size = Pt(22)
    document.styles["Title"].font.bold = True
    document.styles["Quote"].font.italic = True
    document.styles["Quote"].paragraph_format.left_indent = Mm(8)
    document.add_heading(title, 0)

    def inline(paragraph, nodes, marks=None):
        marks = marks or set()
        for node in nodes:
            if isinstance(node, str):
                run = paragraph.add_run(node)
                run.bold = bool(marks & {"b", "strong"})
                run.italic = bool(marks & {"i", "em"})
                run.underline = "u" in marks
                run.font.strike = bool(marks & {"s", "strike"})
                if marks & {"code", "pre"}:
                    run.font.name = "Courier New"
            elif node.tag == "br":
                paragraph.add_run().add_break()
            elif node.tag == "a" and "href" in node.attrs:
                start = len(paragraph._p)
                inline(paragraph, node.children, marks)
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), paragraph.part.relate_to(node.attrs["href"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                for child in list(paragraph._p)[start:]:
                    link.append(child)
                paragraph._p.append(link)
            else:
                inline(paragraph, node.children, marks | {node.tag})

    def numbering(kind, depth, start):
        definitions = document.part.numbering_part.element
        abstract_id = max([int(item.get(qn("w:abstractNumId"))) for item in definitions.findall(qn("w:abstractNum"))], default=-1) + 1
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        for tag, value in (("start", str(start)), ("numFmt", "decimal" if kind == "ol" else "bullet"), ("lvlText", "%1." if kind == "ol" else "•"), ("lvlJc", "left")):
            element = OxmlElement("w:" + tag)
            element.set(qn("w:val"), value)
            level.append(element)
        abstract.append(level)
        definitions.append(abstract)
        return definitions.add_num(abstract_id).numId

    def blocks(nodes, depth=0, list_kind=None, number_id=None):
        for node in nodes:
            if isinstance(node, str):
                if node.strip():
                    inline(document.add_paragraph(), [node])
            elif node.tag in {"ul", "ol"}:
                blocks(node.children, depth + 1, node.tag, numbering(node.tag, depth, int(node.attrs.get("start", 1))))
            elif node.tag == "li":
                paragraph = document.add_paragraph(style="List Number" if list_kind == "ol" else "List Bullet")
                paragraph.paragraph_format.left_indent = Mm(6 * depth)
                paragraph.paragraph_format.first_line_indent = Mm(-3)
                if number_id is not None:
                    props = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                    props.get_or_add_ilvl().val = 0
                    props.get_or_add_numId().val = number_id
                inline(paragraph, [child for child in node.children if not isinstance(child, Node) or child.tag not in {"ul", "ol"}])
                blocks([child for child in node.children if isinstance(child, Node) and child.tag in {"ul", "ol"}], depth)
            elif node.tag == "blockquote":
                for child in node.children:
                    paragraph = document.add_paragraph(style="Quote")
                    inline(paragraph, child.children if isinstance(child, Node) and child.tag == "p" else [child])
            elif node.tag == "hr":
                document.add_paragraph("────────────────────")
            else:
                paragraph = document.add_heading(level=int(node.tag[1])) if node.tag.startswith("h") and node.tag[1:].isdigit() else document.add_paragraph()
                inline(paragraph, node.children, {node.tag})

    blocks(root.children)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
